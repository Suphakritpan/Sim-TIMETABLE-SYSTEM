"""The two downloads on page 6. Excel is checked by a program, so its format follows the brief exactly."""

import io
from html import escape
from pathlib import Path

import openpyxl
from weasyprint import HTML

from app.engine.domain import BLOCKED, DAY_NAMES, DAYS, KIND_NAMES, PERIODS

SHEET_NAME = "ตารางสอน"
HEADERS = ["รหัสรายการสอน", "ประเภท", "วัน", "คาบเริ่ม", "จำนวนคาบ", "รหัสห้อง"]
FONTS = Path(__file__).resolve().parents[1] / "fonts"
LUNCH_ROW = 4  # the lunch row is drawn after this period


def text_cell(sheet, row: int, column: int, value: str) -> None:
    """Write text that a spreadsheet will never run as a formula (A9)."""
    cell = sheet.cell(row=row, column=column, value=value)
    cell.data_type = "s"


def to_excel(items: list[dict]) -> bytes:
    """One sheet, one header row, one row per placed block. Period numbers are numbers; the rest is text."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = SHEET_NAME
    sheet.append(HEADERS)
    for number, item in enumerate(items, start=2):
        text_cell(sheet, number, 1, item["assignmentId"])
        text_cell(sheet, number, 2, KIND_NAMES[item["kind"]])
        text_cell(sheet, number, 3, DAY_NAMES[item["day"]])
        sheet.cell(row=number, column=4, value=item["period"])
        sheet.cell(row=number, column=5, value=item["periodCount"])
        text_cell(sheet, number, 6, item["roomId"])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def cell_text(item: dict, assignment: dict, course: dict) -> str:
    name, kind = escape(course["name"]), KIND_NAMES[item["kind"]][0]
    return f"{name} ({kind})<br>{escape(item['roomId'])} · {escape(assignment['teacherId'])}"


def group_grid(group: dict, items: list[dict], assignments: dict, courses: dict) -> str:
    """One group's week as an HTML table. Two sessions in a cell would both be shown."""
    cells: dict = {}
    for item in items:
        assignment = assignments[item["assignmentId"]]
        if group["id"] in assignment["groupIds"]:
            for p in range(item["period"], item["period"] + item["periodCount"]):
                text = cell_text(item, assignment, courses[assignment["courseCode"]])
                cells.setdefault((item["day"], p), []).append(text)
    blocked = {(b["day"], p) for b in BLOCKED for p in range(b["fromPeriod"], b["toPeriod"] + 1)}
    head = "".join(f"<th>{DAY_NAMES[d]}</th>" for d in DAYS)
    rows = []
    for period in PERIODS:
        boxes = []
        for day in DAYS:
            text = "กิจกรรม" if (day, period) in blocked else "<hr>".join(cells.get((day, period), []))
            boxes.append(f"<td>{text}</td>")
        rows.append(f"<tr><th>{period}</th>{''.join(boxes)}</tr>")
        if period == LUNCH_ROW:
            rows.append(f'<tr class="lunch"><td colspan="{len(DAYS) + 1}">พักกลางวัน 12.00-13.00 น.</td></tr>')
    title = f"{escape(group['name'])} ({escape(group['id'])})"
    return f"<section><h2>{title}</h2><table><tr><th>คาบ</th>{head}</tr>{''.join(rows)}</table></section>"


def to_pdf(items: list[dict], data: dict) -> bytes:
    """A printable A4 timetable: one landscape page per group."""
    style = f"""
    @font-face {{ font-family: Sarabun; src: url('{(FONTS / "Sarabun-Regular.ttf").as_uri()}'); }}
    @font-face {{ font-family: Sarabun; font-weight: bold; src: url('{(FONTS / "Sarabun-Bold.ttf").as_uri()}'); }}
    @page {{ size: A4 landscape; margin: 10mm; }}
    body {{ font-family: Sarabun; font-size: 9pt; }}
    section {{ break-after: page; }}
    h2 {{ font-size: 14pt; margin: 0 0 4mm; }}
    table {{ width: 100%; border-collapse: collapse; table-layout: fixed; }}
    th, td {{ border: 0.5pt solid #555; padding: 1mm; text-align: center; height: 15mm; }}
    th {{ background: #eee; height: auto; }}
    th:first-child {{ width: 10mm; }}
    .lunch td {{ height: auto; background: #eee; }}
    hr {{ border: 0; border-top: 0.5pt dashed #555; }}
    """
    assignments = {a["id"]: a for a in data["assignments"]}
    courses = {c["code"]: c for c in data["courses"]}
    body = "".join(group_grid(g, items, assignments, courses) for g in data["groups"])
    return HTML(
        string=f'<html lang="th"><head><meta charset="utf-8"><style>{style}</style></head><body>{body}</body></html>'
    ).write_pdf()
