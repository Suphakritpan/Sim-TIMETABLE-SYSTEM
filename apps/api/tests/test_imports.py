"""F2: upload and check data. Every planted problem is reported at once, with where and how to fix it."""

import io
import os
import zipfile
from datetime import timedelta
from pathlib import Path

import pytest
from sqlalchemy import func, select, update
from workbooks import base, build, upload

from app import worker
from app.db import SessionLocal
from app.engine.excel_reader import MAX_UNZIPPED_MB
from app.models import AuditLog, Dataset

ROOT = Path(__file__).resolve().parents[3]
LOCAL_FILES = [
    Path.home() / "Desktop/full_test_data/2_ข้อมูลชุดทดลอง_4กลุ่ม.xlsx",
    Path.home() / "Desktop/full_test_data/3_ข้อมูลชุดหลัก_12กลุ่ม.xlsx",
    *sorted((ROOT / "apps/web/public/samples").glob("sample-*.xlsx")),
]


def codes(body: dict) -> set[str]:
    return {i["code"] for i in body["issues"]}


def test_valid_file_gives_summary_and_all_data(client):
    r = upload(client, build())
    body = r.json()
    assert r.status_code == 201 and body["status"] == "valid" and body["issues"] == []
    assert body["summary"] == {"teachers": 2, "rooms": 2, "groups": 2, "courses": 2, "assignments": 2, "requests": 1}
    assert body["importId"].startswith("imp_") and len(body["importId"]) == 36
    assert body["teachers"][0]["unavailable"] == [{"day": "TUE", "fromPeriod": 8, "toPeriod": 8}]
    assert body["assignments"][1]["groupIds"] == ["G1", "G2"]  # spaces after the comma are trimmed
    assert body["periods"]["blocked"] == [{"day": "WED", "fromPeriod": 7, "toPeriod": 8}]
    assert client.get(f"/api/v1/imports/{body['importId']}").json() == body


def test_every_planted_problem_is_listed_at_once(client):
    rows = base()
    rows["ครู"].append(["T1", "ครูซ้ำ", "ไอที", 0])
    rows["เวลาไม่สะดวก"] += [["T9", "จันทร์", 1, 2], ["T1", "เสาร์", 9, 2]]
    rows["ห้อง"].append(["LR-2", "ห้อง", "ห้องบรรยาย", -5, ""])
    rows["กลุ่มเรียน"].append(["G3", "", "ไอที", 0])
    rows["รายวิชา"] += [["C3", "ผิดรูป", "1/2/3", "ไอที", ""], ["C4", "ไม่มีห้อง", "0-2-1", "ไอที", ""]]
    rows["รายวิชา"].append(["C5", "ไม่มีห้องชนิดนี้", "0-2-1", "ไอที", "ห้องแล็บ"])
    rows["รายการสอน"] += [["S1", "C1", "G1", "T1"], ["S3", "C9", "G9", "T9"]]
    r = upload(client, build(rows))
    body = r.json()
    assert r.status_code == 201 and body["status"] == "invalid"  # B23
    assert {"DUPLICATE_ID", "UNKNOWN_REFERENCE", "INVALID_VALUE", "MISSING_VALUE"} <= codes(body)
    assert len(body["issues"]) >= 12
    for issue in body["issues"]:
        assert issue["sheet"] and issue["message"] and issue["fix"]


def test_duplicate_teacher_points_to_the_row(client):
    rows = base()
    rows["ครู"].append(["T1", "ครูซ้ำ", "ไอที", 6])
    issue = upload(client, build(rows)).json()["issues"][0]
    assert issue["code"] == "DUPLICATE_ID" and issue["sheet"] == "ครู" and issue["row"] == 4
    assert "แถว 2" in issue["message"]


def test_group_that_fits_no_room(client):
    rows = base()
    rows["กลุ่มเรียน"][0][3] = 50
    assert "ROOM_TOO_SMALL" in codes(upload(client, build(rows)).json())


def test_combined_class_counts_all_students(client):
    rows = base()
    rows["กลุ่มเรียน"][1][3] = 20  # 30 + 20 = 50, over the lecture room's 45
    body = upload(client, build(rows)).json()
    assert [i["row"] for i in body["issues"] if i["code"] == "ROOM_TOO_SMALL"] == [3]


def test_missing_sheets_and_columns(client):
    rows = base()
    del rows["ห้อง"]
    content = build(rows)
    body = upload(client, content).json()
    assert [i for i in body["issues"] if i["code"] == "MISSING_SHEET"][0]["sheet"] == "ห้อง"

    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(build()))
    book["ครู"].delete_cols(4)
    out = io.BytesIO()
    book.save(out)
    issues = upload(client, out.getvalue()).json()["issues"]
    assert any(i["code"] == "MISSING_COLUMN" and i["column"] == "สอนได้สูงสุดต่อวัน (คาบ)" for i in issues)


def test_column_order_and_text_numbers_are_accepted(client):
    import openpyxl

    book = openpyxl.Workbook()
    book.remove(book.active)
    source = openpyxl.load_workbook(io.BytesIO(build()))
    for sheet in source.worksheets:
        target = book.create_sheet(sheet.title)
        for row in sheet.iter_rows(values_only=True):
            target.append([" " + str(c) + " " if c is not None else None for c in reversed(row)])
    out = io.BytesIO()
    book.save(out)
    assert upload(client, out.getvalue()).json()["status"] == "valid"


def test_new_room_type_is_accepted(client):
    rows = base()
    rows["ห้อง"].append(["AI-1", "ห้องหุ่นยนต์", "ห้องหุ่นยนต์", 40, ""])
    rows["รายวิชา"][0][4] = "ห้องหุ่นยนต์"
    assert upload(client, build(rows)).json()["status"] == "valid"  # B16


def test_block_longer_than_a_morning(client):
    rows = base()
    rows["รายวิชา"][0][2] = "0-5-2"
    assert "BLOCK_TOO_LONG" in codes(upload(client, build(rows)).json())


def test_block_over_the_teachers_daily_limit(client):
    rows = base()
    rows["ครู"][0][3] = 1
    assert "BLOCK_OVER_TEACHER_LIMIT" in codes(upload(client, build(rows)).json())


def test_teacher_without_a_long_enough_free_stretch(client):
    rows = base()
    # free only in single periods: every second period is taken on every day
    rows["เวลาไม่สะดวก"] = [["T1", d, p, p] for d in ("จันทร์", "อังคาร", "พุธ", "พฤหัสบดี", "ศุกร์") for p in (2, 4, 6, 8)]
    assert "NO_FREE_STRETCH" in codes(upload(client, build(rows)).json())


def test_overloaded_teacher_and_group(client):
    rows = base()
    rows["รายวิชา"].append(["C9", "วิชายาว", "4-4-4", "ไอที", "ห้องคอมพิวเตอร์"])
    rows["รายการสอน"] += [[f"X{i}", "C9", "G1", "T1"] for i in range(5)]
    assert {"TEACHER_OVERLOADED", "GROUP_OVERLOADED"} <= codes(upload(client, build(rows)).json())


def test_empty_template_says_only_that_it_holds_no_data(client):
    body = upload(client, build({name: [] for name in base()})).json()
    assert body["status"] == "invalid" and codes(body) == {"EMPTY_SHEET"}


@pytest.mark.parametrize("name", ["data.csv", "data.xls", "data"])
def test_wrong_file_type_is_refused(client, name):
    r = upload(client, b"a,b", name)
    assert r.status_code == 415 and r.json()["error"]["code"] == "UNSUPPORTED_FILE_TYPE"


def test_fake_excel_is_refused(client):
    r = upload(client, b"this is not a workbook")
    assert r.status_code == 400 and r.json()["error"]["code"] == "FILE_UNREADABLE"


def test_zip_that_is_not_excel_is_refused(client):
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("note.txt", "hello")
    assert upload(client, out.getvalue()).status_code == 400


def test_file_over_the_limit_is_refused(client):
    r = upload(client, b"x" * (1024 * 1024 + 1))  # tests run with MAX_UPLOAD_MB=1
    assert r.status_code == 413 and r.json()["error"]["code"] == "FILE_TOO_LARGE"


def test_small_file_that_expands_hugely_is_refused(client):
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/workbook.xml", "<a/>")
        archive.writestr("xl/huge.bin", b"\0" * ((MAX_UNZIPPED_MB + 1) * 1024 * 1024))
    r = upload(client, out.getvalue())
    assert r.status_code == 413 and r.json()["error"]["code"] == "FILE_TOO_LARGE"


def test_wrong_id_gets_404(client):
    r = client.get("/api/v1/imports/imp_nothing")
    assert r.status_code == 404 and r.json()["error"]["code"] == "IMPORT_NOT_FOUND"


def test_upload_is_audited_without_names(client):
    upload(client, build())
    with SessionLocal() as db:
        row = db.scalars(select(AuditLog)).one()
    assert row.action == "upload" and "ครู" not in str(row.detail)


def test_cleanup_deletes_old_imports_only(client):
    old = upload(client, build()).json()["importId"]
    new = upload(client, build()).json()["importId"]
    with SessionLocal() as db:
        db.execute(
            update(Dataset)
            .where(Dataset.public_id == old)
            .values(created_at=func.now() - timedelta(days=worker.DATASET_DAYS + 1))
        )
        db.commit()
    worker.cleanup()
    assert client.get(f"/api/v1/imports/{old}").status_code == 404
    assert client.get(f"/api/v1/imports/{new}").status_code == 200


@pytest.mark.parametrize("path", [p for p in LOCAL_FILES if p.exists()], ids=lambda p: p.name)
def test_provided_files_are_valid(client, path, monkeypatch):
    # local only: these files are not always on the machine
    monkeypatch.setattr("app.config.MAX_UPLOAD_MB", 10)
    body = upload(client, path.read_bytes(), path.name).json()
    assert body["status"] == "valid", body["issues"][:5]
    assert body["summary"]["assignments"] > 0
    assert os.path.exists(path)
