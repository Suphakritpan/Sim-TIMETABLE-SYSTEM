"""F7: the downloads. Excel is read back and compared with the format in the brief."""

import io
import shutil
import subprocess
from pathlib import Path

import openpyxl
import pytest
from workbooks import build, upload

from app import worker
from app.engine.exports import HEADERS, SHEET_NAME, to_excel

SAMPLE = Path.home() / "Desktop/full_test_data/4_ตัวอย่างไฟล์ตารางสอนที่ต้องส่งออก.xlsx"


def finished_job(client, seed=5) -> str:
    import_id = upload(client, build()).json()["importId"]
    job = client.post("/api/v1/schedules", json={"importId": import_id, "seed": seed}).json()
    worker.run_one()
    return job["jobId"]


def read_back(content: bytes) -> list[tuple]:
    book = openpyxl.load_workbook(io.BytesIO(content))
    assert book.sheetnames == [SHEET_NAME]
    return list(book[SHEET_NAME].iter_rows(values_only=True))


def test_excel_matches_the_format(client):
    r = client.get(f"/api/v1/schedules/{finished_job(client)}/export", params={"format": "xlsx"})
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    assert r.headers["content-disposition"].startswith("attachment; filename=")
    rows = read_back(r.content)
    assert list(rows[0]) == HEADERS and len(rows) == 1 + 3  # S1 has 2 blocks, S2 has 1
    for row in rows[1:]:
        assert row[1] in ("ทฤษฎี", "ปฏิบัติ") and row[2] in ("จันทร์", "อังคาร", "พุธ", "พฤหัสบดี", "ศุกร์")
        assert (
            isinstance(row[3], int) and isinstance(row[4], int) and isinstance(row[0], str) and isinstance(row[5], str)
        )


@pytest.mark.skipif(not SAMPLE.exists(), reason="competition files are not on this machine")
def test_headers_and_types_equal_the_official_sample():
    official = openpyxl.load_workbook(SAMPLE)
    item = {"assignmentId": "S0007", "kind": "theory", "day": "MON", "period": 1, "periodCount": 1, "roomId": "LR-02"}
    ours = read_back(to_excel([item]))
    sheet = official[SHEET_NAME]
    rows = list(sheet.iter_rows(values_only=True))
    assert official.sheetnames == [SHEET_NAME] and list(rows[0]) == list(ours[0])
    assert ours[1] == rows[1]  # the same teaching item gives the same cells, with the same types


def test_same_seed_gives_the_same_cells(client):
    first = client.get(f"/api/v1/schedules/{finished_job(client)}/export", params={"format": "xlsx"})
    second = client.get(f"/api/v1/schedules/{finished_job(client)}/export", params={"format": "xlsx"})
    assert read_back(first.content) == read_back(second.content)


def test_text_is_never_a_formula():
    item = {"assignmentId": "=1+1", "kind": "theory", "day": "MON", "period": 1, "periodCount": 1, "roomId": "=A1"}
    book = openpyxl.load_workbook(io.BytesIO(to_excel([item])), data_only=False)
    cell = book[SHEET_NAME]["A2"]
    assert cell.data_type == "s" and cell.value == "=1+1"


def test_export_of_an_unfinished_job_waits(client):
    import_id = upload(client, build()).json()["importId"]
    job = client.post("/api/v1/schedules", json={"importId": import_id}).json()
    r = client.get(f"/api/v1/schedules/{job['jobId']}/export", params={"format": "xlsx"})
    assert r.status_code == 409 and r.json()["error"]["code"] == "SCHEDULE_NOT_READY"


def test_unknown_format_and_id(client):
    job = finished_job(client)
    r = client.get(f"/api/v1/schedules/{job}/export", params={"format": "csv"})
    assert r.status_code == 400 and r.json()["error"]["code"] == "EXPORT_FORMAT_UNSUPPORTED"
    assert client.get("/api/v1/schedules/nope/export", params={"format": "xlsx"}).status_code == 404


def test_pdf_opens_and_embeds_thai(client, tmp_path):
    r = client.get(f"/api/v1/schedules/{finished_job(client)}/export", params={"format": "pdf"})
    assert r.status_code == 200 and r.content.startswith(b"%PDF") and r.headers["content-type"] == "application/pdf"
    if shutil.which("pdftotext"):
        path = tmp_path / "timetable.pdf"
        path.write_bytes(r.content)
        text = subprocess.run(["pdftotext", str(path), "-"], capture_output=True, text=True, check=True).stdout
        fonts = subprocess.run(["pdffonts", str(path)], capture_output=True, text=True, check=True).stdout
        assert "วิชาหนึ่ง" in text and "พักกลางวัน" in text and "Sarabun" in fonts
